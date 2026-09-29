"""End-to-end sync with every provider mocked (no network)."""

import csv
import json

import httpx
import openpyxl
import pytest

from citereclaim.models import (
    Action,
    ScopusArticleStatus,
    TargetVersion,
    TrackedPaper,
)
from citereclaim.pipeline import Providers, Syncer
from citereclaim.providers import scopus_sources
from citereclaim.reporting import export_support, report_json

from .conftest import load_fixture

VOR = "10.1016/j.media.2026.103932"
SCOPUS_URL = "https://api.elsevier.com/content/search/scopus"


def load_sources(db, tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scopus Sources Aug. 2026"
    ws.append(
        ["Sourcerecord ID", "Source Title", "ISSN", "EISSN", "Active or Inactive", "Coverage"]
    )
    ws.append([1, "Medical Image Analysis", "13618415", "", "Active", "1996-2026"])
    ws.append([2, "Diagnostics", "", "20754418", "Active", "2011-2026"])
    ws.append(
        [3, "Pattern Analysis and Applications", "14337541", "1433755X", "Active", "1998-2026"]
    )
    p = tmp_path / "ext_list_Aug_2026.xlsx"
    wb.save(p)
    scopus_sources.import_file(db, p)


def make_router(openalex_down=False, scopus=None):
    fixtures = {
        f"https://api.crossref.org/works/{VOR}": load_fixture("crossref_vor.json"),
        "https://api.crossref.org/works/10.3390/diagnostics16152385": load_fixture(
            "crossref_citing_mixed.json"
        ),
        "https://api.crossref.org/works/10.1007/s10044-026-01765-1": load_fixture(
            "crossref_citing_vor.json"
        ),
        "https://api.crossref.org/works/10.9999/thesis.1": load_fixture(
            "crossref_citing_noref.json"
        ),
        f"https://api.semanticscholar.org/graph/v1/paper/DOI:{VOR}": load_fixture("s2_paper.json"),
        "https://api.semanticscholar.org/graph/v1/paper/ARXIV:2504.08635": load_fixture(
            "s2_paper.json"
        ),
        "https://api.semanticscholar.org/graph/v1/paper/S2PAPER/citations": load_fixture(
            "s2_citations.json"
        ),
        f"https://api.openalex.org/works/doi:{VOR}": load_fixture("openalex_vor.json"),
    }
    calls: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url.copy_with(query=None))
        calls.append(str(request.url))
        if url.startswith("https://api.openalex.org") and openalex_down:
            return httpx.Response(500)
        if url.startswith("https://api.elsevier.com"):
            return scopus(request) if scopus else httpx.Response(401)
        if url in fixtures:
            return httpx.Response(200, json=fixtures[url])
        if url == "https://api.openalex.org/works":
            flt = request.url.params.get("filter", "")
            if flt == "cites:W100":
                return httpx.Response(200, json=load_fixture("openalex_cites.json"))
            return httpx.Response(200, json={"meta": {"count": 0}, "results": []})
        if url == "https://api.crossref.org/works":
            return httpx.Response(200, json={"message": {"items": []}})
        return httpx.Response(404)

    return handler, calls


def run_sync(settings, db, handler):
    db.upsert_paper(
        TrackedPaper(
            name="DEMO",
            arxiv_id="2504.08635",
            arxiv_doi="10.48550/arxiv.2504.08635",
            journal_doi=VOR,
        )
    )
    paper = db.get_paper("DEMO")
    transport = httpx.MockTransport(handler)
    providers = Providers.build(settings, db, transport=transport, sleep=lambda s: None)
    try:
        return Syncer(settings, db, providers).sync(paper), db.get_paper("DEMO")
    finally:
        providers.close()


def by_doi(db, paper):
    return {c["work"].ids.doi: c for c in db.citations_for_paper(paper.id)}


def test_sync_without_any_keys(settings, db, tmp_path):
    load_sources(db, tmp_path)
    handler, calls = make_router()
    rep, paper = run_sync(settings, db, handler)
    assert rep.status == "ok"
    cites = by_doi(db, paper)
    assert set(cites) == {
        "10.3390/diagnostics16152385",
        "10.1007/s10044-026-01765-1",
        "10.9999/thesis.1",
        "10.48550/arxiv.2601.14584",
    }
    powdr = cites["10.3390/diagnostics16152385"]
    # Found in two graphs, deduplicated by DOI.
    assert powdr["sources"] == ["openalex", "semantic_scholar"]
    assert powdr["result"].target_version == TargetVersion.PREPRINT
    assert powdr["result"].scopus_article_status == ScopusArticleStatus.SCOPUS_SOURCE_ONLY
    assert powdr["result"].action == Action.CHECK_SCOPUS
    vor = cites["10.1007/s10044-026-01765-1"]["result"]
    assert vor.target_version == TargetVersion.VERSION_OF_RECORD and vor.action == Action.OK
    thesis = cites["10.9999/thesis.1"]["result"]
    assert thesis.action == Action.NON_SCOPUS and thesis.target_version == TargetVersion.UNKNOWN
    assert cites["10.48550/arxiv.2601.14584"]["result"].action == Action.PREPRINT_ONLY
    # Paper metadata was enriched from Crossref.
    assert paper.journal_title == "Medical Image Analysis"
    assert paper.version_ids["vor"]["openalex"] == "W100"
    # Scopus API never called without a key.
    assert not any("elsevier" in c for c in calls)

    # Report JSON and support export.
    data = report_json(db, paper)
    assert data["summary"]["unique_citing_works"] == 4
    assert data["summary"]["likely_citing_preprint"] == 1
    files, n, missing = export_support(db, paper, tmp_path / "out")
    xlsx_path, txt_path, csv_path, md_path = files
    assert n == 1
    assert missing == 2  # no API key: cited and citing Scopus links unknown
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["Reference linking"]
    header = [c.value for c in ws[1]]
    assert header == [
        "Cited article title",
        "Cited article link in Scopus",
        "Citing article",
        "Citing article link in Scopus",
        "Final DOI",
        "Preprint / arXiv",
        "Reference / evidence",
        "Requested correction",
    ]
    row = [c.value for c in ws[2]]
    assert row[2].startswith("POWDR")
    assert row[4] == VOR and row[5] == "arXiv:2504.08635"
    assert row[6].startswith("Reference cites arXiv:2504.08635")
    assert row[7].startswith("Link this reference to the Scopus record")
    assert not any(c.hyperlink for r in ws.iter_rows() for c in r)  # see write_linking_xlsx
    txt = txt_path.read_text()
    assert 'Missing citation links to published version of "Latent diffusion autoencoders"' in txt
    assert "[PASTE THE SCOPUS LINK OF THE PUBLISHED ARTICLE]" in txt
    rows = list(csv.DictReader(csv_path.open()))
    assert rows[0]["citing_doi"] == "10.3390/diagnostics16152385"
    assert rows[0]["final_doi"] == VOR and "arXiv" in rows[0]["preprint_evidence"]
    assert "POWDR" in md_path.read_text()

    # A second sync is served from the cache.
    n_calls = len(calls)
    run_sync(settings, db, handler)
    assert len(calls) == n_calls


def test_sync_with_scopus_key_detects_missing_link(settings, db, tmp_path):
    load_sources(db, tmp_path)
    settings.elsevier_api_key = "KEY"

    def entry(eid, doi):
        return {
            "eid": eid,
            "prism:doi": doi,
            "dc:title": "t",
            "prism:coverDate": "2026-01-01",
            "subtype": "ar",
        }

    def scopus(request):
        q = request.url.params["query"]
        if q == f'DOI("{VOR}")':
            res = [entry("2-s2.0-VOR", VOR)]
        elif q == "REFEID(2-s2.0-VOR)":
            res = [entry("2-s2.0-GEN", "10.1007/s10044-026-01765-1")]
        elif q == 'DOI("10.3390/diagnostics16152385")':
            res = [
                {
                    **entry("2-s2.0-POWDR", "10.3390/diagnostics16152385"),
                    "link": [{"@ref": "scopus", "@href": "https://scopus.example/POWDR"}],
                }
            ]
        else:
            return httpx.Response(
                200,
                json={
                    "search-results": {
                        "opensearch:totalResults": "0",
                        "entry": [{"error": "Result set was empty"}],
                    }
                },
            )
        return httpx.Response(
            200, json={"search-results": {"opensearch:totalResults": str(len(res)), "entry": res}}
        )

    handler, _ = make_router(scopus=scopus)
    _, paper = run_sync(settings, db, handler)
    cites = by_doi(db, paper)
    powdr = cites["10.3390/diagnostics16152385"]
    assert powdr["result"].scopus_article_status == ScopusArticleStatus.SCOPUS_CONFIRMED
    assert powdr["result"].action == Action.LIKELY_MISSING_LINK
    assert powdr["work"].ids.scopus_eid == "2-s2.0-POWDR"
    gen = cites["10.1007/s10044-026-01765-1"]
    assert gen["result"].action == Action.OK
    files, _, missing = export_support(db, paper, tmp_path / "out2")
    ws = openpyxl.load_workbook(files[0])["Reference linking"]
    row = [c.value for c in ws[2]]
    assert row[1] == "https://www.scopus.com/record/display.uri?eid=2-s2.0-VOR&origin=resultslist"
    assert row[3] == "https://scopus.example/POWDR"  # link returned by the API
    assert missing == 0
    assert "scopus" in gen["sources"]
    thesis = cites["10.9999/thesis.1"]["result"]
    assert thesis.scopus_article_status == ScopusArticleStatus.SCOPUS_NOT_FOUND


def test_scopus_permission_denied_does_not_break_sync(settings, db, tmp_path):
    load_sources(db, tmp_path)
    settings.elsevier_api_key = "BAD"
    handler, _ = make_router(scopus=lambda r: httpx.Response(403))
    rep, paper = run_sync(settings, db, handler)
    assert rep.status == "ok"
    powdr = by_doi(db, paper)["10.3390/diagnostics16152385"]["result"]
    assert powdr.scopus_article_status == ScopusArticleStatus.SCOPUS_PERMISSION_DENIED
    assert powdr.action == Action.CHECK_SCOPUS  # not NON_SCOPUS: denial ≠ not indexed
    assert any("Scopus VoR lookup" in w for w in rep.warnings)


def test_openalex_unavailable_keeps_going_and_does_not_prune(settings, db, tmp_path):
    load_sources(db, tmp_path)
    handler, _ = make_router()
    run_sync(settings, db, handler)
    assert len(db.citations_for_paper(1)) == 4
    db.cache_clear()
    down, _ = make_router(openalex_down=True)
    rep, paper = run_sync(settings, db, down)
    assert rep.status == "partial"
    assert any("openalex" in w for w in rep.warnings)
    # The thesis/Generative records were only in OpenAlex; they must not be pruned.
    assert len(db.citations_for_paper(paper.id)) == 4


def test_scholar_import_is_merged(settings, db, tmp_path):
    from citereclaim.providers.scholar_import import parse_file, row_hash

    from .conftest import FIXTURES

    load_sources(db, tmp_path)
    handler, _ = make_router()
    run_sync(settings, db, handler)
    rows = parse_file(FIXTURES / "scholar_citations.csv")
    db.add_scholar_rows(1, "x.csv", [(row_hash(r), r) for r in rows])
    _, paper = run_sync(settings, db, handler)
    cites = db.citations_for_paper(paper.id)
    powdr = next(c for c in cites if c["work"].ids.doi == "10.3390/diagnostics16152385")
    # POWDR had no DOI in the Scholar file; it was merged by title with the graph records.
    assert "google_scholar_import" in powdr["sources"]
    unrelated = [c for c in cites if c["work"].title and "unrelated" in c["work"].title]
    assert len(unrelated) == 1
    assert unrelated[0]["sources"] == ["google_scholar_import"]
    # Scholar-only records never get article-level Scopus claims.
    assert unrelated[0]["result"].scopus_article_status != ScopusArticleStatus.SCOPUS_CONFIRMED
    assert json.dumps(report_json(db, paper))  # serialisable


@pytest.mark.parametrize("missing", ["journal", "arxiv"])
def test_paper_with_single_version(settings, db, tmp_path, missing):
    load_sources(db, tmp_path)
    handler, _ = make_router()
    kwargs = (
        {"arxiv_id": "2504.08635", "arxiv_doi": "10.48550/arxiv.2504.08635"}
        if missing == "journal"
        else {"journal_doi": VOR}
    )
    db.upsert_paper(TrackedPaper(name="ONE", **kwargs))
    providers = Providers.build(
        settings, db, transport=httpx.MockTransport(handler), sleep=lambda s: None
    )
    rep = Syncer(settings, db, providers).sync(db.get_paper("ONE"))
    providers.close()
    assert rep.stats["unique_works"] >= 2
    if missing == "journal":
        assert db.get_paper("ONE").journal_doi == VOR  # inferred from Semantic Scholar
        assert any("inferred from Semantic Scholar" in w for w in rep.warnings)


def test_manual_citation_is_included_with_confirmed_target(settings, db, tmp_path):
    load_sources(db, tmp_path)
    handler, _ = make_router()
    run_sync(settings, db, handler)
    # The thesis reference list has no usable entry: confirm by hand it cites the preprint.
    db.add_manual_citation(1, "10.9999/thesis.1", "PREPRINT", "seen in Scopus")
    db.add_manual_citation(1, "10.3390/diagnostics16152385", "PREPRINT", None)
    _, paper = run_sync(settings, db, handler)
    cites = by_doi(db, paper)
    thesis = cites["10.9999/thesis.1"]
    assert "manual" in thesis["sources"]
    assert thesis["result"].target_version == TargetVersion.PREPRINT
    assert "confirmed manually" in thesis["result"].target_evidence[0]
    assert any(e.signal == "manual_confirmation" for e in thesis["result"].confidence_evidence)
    # A manual entry for an already-discovered work merges instead of duplicating.
    assert len(cites) == 4
    files, _, _ = export_support(db, paper, tmp_path / "out")
    rows = list(openpyxl.load_workbook(files[0]).active.iter_rows(min_row=2, values_only=True))
    assert all(r[2] != thesis["work"].title for r in rows)  # NON_SCOPUS is never exported
    manual_row = next(r for r in rows if r[2].startswith("POWDR"))
    assert manual_row[6].startswith(
        "Reference cites arXiv:2504.08635 instead of the published record."
    )


def test_citedby_count_deduces_missing_link_without_refeid(settings, db, tmp_path):
    """REFEID not entitled, but the VoR cited-by count proves a link is missing."""
    load_sources(db, tmp_path)
    settings.elsevier_api_key = "KEY"
    restricted = {
        "service-error": {
            "status": {
                "statusText": "Use of certain field restrictions in the search query is "
                "not allowed for this requestor."
            }
        }
    }

    def entry(eid, doi, **extra):
        return {
            "eid": eid,
            "prism:doi": doi,
            "dc:title": "t",
            "prism:coverDate": "2026-01-01",
            "subtype": "ar",
            **extra,
        }

    def scopus(request):
        if "/content/abstract/" in request.url.path:
            return httpx.Response(401)  # view=REF not entitled off-campus
        q = request.url.params["query"]
        if q.startswith("REFEID"):
            return httpx.Response(400, json=restricted)
        hits = {
            f'DOI("{VOR}")': entry(
                "2-s2.0-VOR",
                VOR,
                **{
                    "citedby-count": "1",
                    "link": [{"@ref": "scopus-citedby", "@href": "https://scopus.example/citedby"}],
                },
            ),
            'DOI("10.3390/diagnostics16152385")': entry("2-s2.0-P", "10.3390/diagnostics16152385"),
            'DOI("10.1007/s10044-026-01765-1")': entry("2-s2.0-G", "10.1007/s10044-026-01765-1"),
        }
        res = [hits[q]] if q in hits else []
        return httpx.Response(
            200,
            json={
                "search-results": {
                    "opensearch:totalResults": str(len(res)),
                    "entry": res or [{"error": "empty"}],
                }
            },
        )

    handler, _ = make_router(scopus=scopus)
    rep, paper = run_sync(settings, db, handler)
    assert any("REFEID() is not permitted" in w for w in rep.warnings)
    cites = by_doi(db, paper)
    powdr = cites["10.3390/diagnostics16152385"]["result"]
    # 2 confirmed citers, Scopus counts 1, and that 1 is explained by the explicit final-DOI
    # citation -> the preprint citation cannot be counted.
    assert powdr.action == Action.LIKELY_MISSING_LINK
    assert any("at least 1 citation(s) are not linked" in e for e in powdr.source_evidence)
    assert any(e.signal == "linkage_deduced_from_count" for e in powdr.confidence_evidence)
    assert cites["10.1007/s10044-026-01765-1"]["result"].action == Action.OK
    meta = db.get_meta(f"scopus_vor:{paper.id}")
    assert meta["citedby_count"] == 1 and meta["deficit"] == 1
    files, _, _ = export_support(db, paper, tmp_path / "out")
    txt = files[1].read_text()
    assert "Scopus currently shows 1 citation for" in txt and "scopus.example/citedby" in txt
    assert "additional Scopus-indexed articles" in txt and "at least" not in txt


def test_scopus_reference_list_proves_missing_link(settings, db, tmp_path):
    """view=REF shows the reference linked to the preprint's Scopus record, not the VoR."""
    load_sources(db, tmp_path)
    settings.elsevier_api_key = "KEY"

    def entry(eid, doi):
        return {
            "eid": eid,
            "prism:doi": doi,
            "dc:title": "t",
            "prism:coverDate": "2026-01-01",
            "subtype": "ar",
        }

    def ref(pos, eid, title, source, ref_type, doi=None):
        return {
            "@id": pos,
            "scopus-eid": eid,
            "type": ref_type,
            "title": title,
            "sourcetitle": source,
            "ce:doi": doi,
            "author-list": {"author": [{"ce:surname": "Lozupone"}]},
        }

    title = (
        "Latent Diffusion Autoencoders: Toward Efficient and Meaningful Unsupervised "
        "Representation Learning in Medical Imaging"
    )
    reference_lists = {
        "2-s2.0-P": [
            ref("1", "2-s2.0-X", "Other paper", "arXiv", "originalReference/other"),
            ref("6", "2-s2.0-ARXIV", title, "arXiv", "originalReference/other"),
        ],
        "2-s2.0-G": [
            ref("3", "2-s2.0-VOR", title, "Medical Image Analysis", "resolvedReference", VOR)
        ],
    }

    def scopus(request):
        path = request.url.path
        if "/content/abstract/eid/" in path:
            eid = path.rsplit("/", 1)[-1]
            refs = reference_lists.get(eid, [])
            return httpx.Response(
                200,
                json={
                    "abstracts-retrieval-response": {
                        "references": {"@total-references": str(len(refs)), "reference": refs}
                    }
                },
            )
        q = request.url.params["query"]
        if q.startswith("REFEID"):
            return httpx.Response(400, json={"x": "not allowed for this requestor"})
        hits = {
            f'DOI("{VOR}")': entry("2-s2.0-VOR", VOR),
            'DOI("10.3390/diagnostics16152385")': entry("2-s2.0-P", "10.3390/diagnostics16152385"),
            'DOI("10.1007/s10044-026-01765-1")': entry("2-s2.0-G", "10.1007/s10044-026-01765-1"),
        }
        res = [hits[q]] if q in hits else []
        return httpx.Response(
            200,
            json={
                "search-results": {
                    "opensearch:totalResults": str(len(res)),
                    "entry": res or [{"error": "empty"}],
                }
            },
        )

    handler, _ = make_router(scopus=scopus)
    _, paper = run_sync(settings, db, handler)
    cites = by_doi(db, paper)
    powdr = cites["10.3390/diagnostics16152385"]
    r = powdr["result"]
    assert r.linkage_status.value == "NOT_LINKED"
    assert r.action == Action.LIKELY_MISSING_LINK
    assert any(
        "linked to record 2-s2.0-ARXIV" in e and "not to the Version of Record" in e
        for e in r.source_evidence
    )
    assert powdr["work"].extra["scopus_linked_eid"] == "2-s2.0-ARXIV"
    assert any(e.signal == "linkage_verified" for e in r.confidence_evidence)
    gen = cites["10.1007/s10044-026-01765-1"]["result"]
    assert gen.linkage_status.value == "LINKED_TO_VOR" and gen.action == Action.OK
    files, _, _ = export_support(db, paper, tmp_path / "out")
    rows = list(openpyxl.load_workbook(files[0]).active.iter_rows(min_row=2, values_only=True))
    powdr_row = next(r for r in rows if r[2].startswith("POWDR"))
    assert powdr_row[6].startswith(
        "Reference cites arXiv:2504.08635 and is currently linked in Scopus to record "
        "2-s2.0-ARXIV instead of published record 2-s2.0-VOR."
    )
    assert (
        powdr_row[7]
        == "Link this reference to Scopus record 2-s2.0-VOR (published Version of Record)."
    )
    txt = files[1].read_text()
    assert "currently linked to record 2-s2.0-ARXIV" in txt and "merge" not in txt
    assert "linked to the published article 2-s2.0-VOR" in txt


def test_support_pack_skips_articles_not_in_scopus(settings, db, tmp_path):
    load_sources(db, tmp_path)
    settings.elsevier_api_key = "KEY"

    # Scopus knows only the VoR: every citing article is SCOPUS_NOT_FOUND.
    def scopus(request):
        if "/content/abstract/" in request.url.path:
            return httpx.Response(401)
        q = request.url.params["query"]
        res = [{"eid": "2-s2.0-VOR", "prism:doi": VOR}] if q == f'DOI("{VOR}")' else []
        return httpx.Response(
            200,
            json={
                "search-results": {
                    "opensearch:totalResults": str(len(res)),
                    "entry": res or [{"error": "empty"}],
                }
            },
        )

    handler, _ = make_router(scopus=scopus)
    _, paper = run_sync(settings, db, handler)
    powdr = by_doi(db, paper)["10.3390/diagnostics16152385"]["result"]
    assert powdr.scopus_article_status == ScopusArticleStatus.SCOPUS_NOT_FOUND
    _, n, _ = export_support(db, paper, tmp_path / "out")
    assert n == 0
