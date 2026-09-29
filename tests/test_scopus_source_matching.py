from pathlib import Path

import openpyxl
import pytest

from citation_reconciler.models import ScopusSourceStatus, WorkRecord
from citation_reconciler.providers import scopus_sources as ss


def build_xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scopus Sources Aug. 2026"
    ws.append(
        [
            "Sourcerecord ID",
            "Source Title",
            "ISSN",
            "EISSN",
            "Active or Inactive",
            "Coverage",
            "Titles Discontinued by Scopus",
            "Source Type",
            "Publisher",
        ]
    )
    ws.append(
        [
            17271,
            "Medical Image Analysis",
            "13618415",
            "13618423",
            "Active",
            "1996-2026",
            "",
            "Journal",
            "Elsevier B.V.",
        ]
    )
    ws.append(
        [21100852989, "Diagnostics", "", "20754418", "Active", "2011-2025", "", "Journal", "MDPI"]
    )
    ws.append(
        [999, "Old Journal of Things", "1234567X", "", "Inactive", "1990-2001", "", "Journal", "X"]
    )
    ws.append([1000, "Titles Only Review", "", "", "Active", "2015-2026", "", "Journal", "Y"])
    ws.append(
        [1001, "Leading Zero Journal", "2349130", "", "Active", "2020-2026", "", "Journal", "Z"]
    )
    acc = wb.create_sheet("Accepted Titles Aug. 2026")
    acc.append(["Status: August 2026", "Accepted titles"])
    acc.append(["Source Title", "ISSN", "EISSN", "Date of Acceptance", "Publisher"])
    acc.append(["Not Yet Indexed Journal", "11112222", "", "2026-08-01", "P"])
    disc = wb.create_sheet("Discontinued Titles Aug. 2026")
    disc.append(["Status (recently discontinued titles in bold)", None, None])
    disc.append(
        [
            "Sourcerecord ID",
            "Source Title",
            "ISSN",
            "EISSN",
            "Publisher",
            "Indexation Change",
            "Year",
        ]
    )
    disc.append([555, "Predatory Journal", "22223333", "", "Q", "Discontinuation", 2018])
    conf = wb.create_sheet("All Conf. Proceedings Jun. 2026")
    conf.append(["Source Title", "Proceedinds Title", "ISBN", "Year", "Volume", "EID"])
    conf.append(
        [
            "2025 28th International Conference on Computer and Information Technology, ICCIT 2025",
            "ICCIT 2025",
            9798331500001,
            2025,
            None,
            1,
        ]
    )
    conf.append(
        [
            "11th International Symposium on Signal, Image, Video and Communications, "
            "ISIVC 2022 - Conference Proceedings",
            "x",
            9781665487245,
            2022,
            None,
            2,
        ]
    )
    out = path / "ext_list_Aug_2026.xlsx"
    wb.save(out)
    return out


@pytest.fixture
def loaded(db, tmp_path):
    ss.import_file(db, build_xlsx(tmp_path))
    return ss.SourceIndex(db)


def test_import_parses_sheets_and_skips_accepted(db, tmp_path):
    res = ss.import_file(db, build_xlsx(tmp_path))
    assert res["version"] == "Aug 2026"
    assert res["sources"] == 6  # 5 main + 1 discontinued; accepted titles skipped
    assert res["isbns"] == 2
    st = ss.status(db)
    assert st["n_sources"] == 6 and st["origin"].endswith(".xlsx") and st["age_days"] == 0


def test_issn_match_and_coverage(loaded):
    m = loaded.match(WorkRecord(issns=["1361-8423"], year=2026))
    assert m.status == ScopusSourceStatus.COVERED and m.matched_on == "issn:1361-8423"
    assert m.source_title == "Medical Image Analysis"


def test_active_source_after_listed_coverage_counts_as_covered(loaded):
    m = loaded.match(WorkRecord(issns=["2075-4418"], year=2026))
    assert m.status == ScopusSourceStatus.COVERED
    assert any("list lag" in e for e in m.evidence)


def test_inactive_source_outside_coverage(loaded):
    m = loaded.match(WorkRecord(issns=["1234-567X"], year=2020))
    assert m.status == ScopusSourceStatus.OUTSIDE_COVERAGE


def test_discontinued_sheet_source(loaded):
    m = loaded.match(WorkRecord(issns=["2222-3333"], year=2024))
    assert m.status == ScopusSourceStatus.OUTSIDE_COVERAGE
    assert any("Discontinu" in e for e in m.evidence)


def test_unknown_issn_is_not_in_list(loaded):
    m = loaded.match(WorkRecord(issns=["1598-8619"], year=2025))
    assert m.status == ScopusSourceStatus.NOT_IN_LIST


def test_accepted_titles_are_not_coverage(loaded):
    assert (
        loaded.match(WorkRecord(issns=["1111-2222"], year=2026)).status
        == ScopusSourceStatus.NOT_IN_LIST
    )


def test_title_fallback_only_without_issn(loaded):
    m = loaded.match(WorkRecord(venue="Titles-Only Review", year=2020))
    assert m.status == ScopusSourceStatus.COVERED_TITLE_MATCH
    # With an ISSN present, the title is never used.
    m2 = loaded.match(WorkRecord(venue="Titles Only Review", issns=["9999-9999"], year=2020))
    assert m2.status == ScopusSourceStatus.NOT_IN_LIST


def test_leading_zero_issn_restored(loaded):
    assert (
        loaded.match(WorkRecord(issns=["0234-9130"], year=2024)).status
        == ScopusSourceStatus.COVERED
    )


def test_isbn_proceedings_match(loaded):
    m = loaded.match(WorkRecord(isbns=["9781665487245"], year=2022))
    assert m.status == ScopusSourceStatus.COVERED and m.matched_on.startswith("isbn")


def test_conference_same_edition_and_series_heuristic(loaded):
    same = loaded.match(
        WorkRecord(
            venue="2025 28th International Conference on Computer and Information "
            "Technology (ICCIT)",
            year=2025,
        )
    )
    assert same.status == ScopusSourceStatus.COVERED_TITLE_MATCH
    later = loaded.match(
        WorkRecord(
            venue="2026 IEEE 13th International Symposium on Signal, Image, Video and "
            "Communications (ISIVC)",
            year=2026,
        )
    )
    assert later.status == ScopusSourceStatus.SERIES_PREVIOUSLY_INDEXED
    assert any("does NOT prove" in e for e in later.evidence)


def test_list_not_loaded(db):
    assert (
        ss.SourceIndex(db).match(WorkRecord(issns=["1361-8415"])).status
        == ScopusSourceStatus.LIST_NOT_LOADED
    )


def test_csv_import(db, tmp_path):
    p = tmp_path / "sources.csv"
    p.write_text("Source Title;Print-ISSN;E-ISSN;Coverage\nSome Journal;1361-8415;;2000-2026\n")
    assert ss.import_file(db, p)["sources"] == 1
    assert (
        ss.SourceIndex(db).match(WorkRecord(issns=["1361-8415"], year=2010)).status
        == ScopusSourceStatus.COVERED
    )


def test_bad_file_rejected(db, tmp_path):
    p = tmp_path / "x.csv"
    p.write_text("foo,bar\n1,2\n")
    with pytest.raises(ValueError):
        ss.import_file(db, p)


def test_parse_coverage():
    assert ss.parse_coverage("2026; 2023-2024; 1995-2005") == [
        (2026, 2026),
        (2023, 2024),
        (1995, 2005),
    ]
    assert ss.parse_coverage("2010-present") == [(2010, 9999)]
    assert ss.parse_coverage(None) == []


def test_discover_download_url_from_contentful_json():
    page = (
        '<a href="//downloads.ctfassets.net/a/b/c/Scopus_book_list_Q2.xlsx">Book list</a>'
        '..."title":"Scopus source title list - August 2026","url":"//downloads.ctfassets.net'
        '/o78/7x/69c/ext_list_Aug_2026.xlsx"...'
    )
    url, label = ss.discover_download_url(page)
    assert url == "https://downloads.ctfassets.net/o78/7x/69c/ext_list_Aug_2026.xlsx"
    assert ss.version_label_from(label) == "Aug 2026"


def test_discover_download_url_from_anchor_and_failure():
    page = (
        '<a class="download" href="/files/list_2026.xlsx" aria-label="Download the Source '
        'title list">x</a>'
    )
    url, _ = ss.discover_download_url(page)
    assert url == "https://www.elsevier.com/files/list_2026.xlsx"
    with pytest.raises(LookupError):
        ss.discover_download_url("<html>nothing</html>")
