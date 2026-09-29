"""Terminal reports, CSV/JSON export and the Scopus-support evidence pack."""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from rich.console import Console, Group
from rich.panel import Panel
from rich.table import Table
from rich.text import Text

from .db import Database
from .models import (
    Action,
    ScopusArticleStatus,
    ScopusSourceStatus,
    TargetVersion,
    TrackedPaper,
)

TARGET_LABEL = {
    TargetVersion.PREPRINT: "arXiv",
    TargetVersion.VERSION_OF_RECORD: "final",
    TargetVersion.BOTH: "both",
    TargetVersion.UNKNOWN: "?",
}
ACTION_STYLE = {
    Action.OK: "green",
    Action.CHECK_SCOPUS: "yellow",
    Action.LIKELY_MISSING_LINK: "bold red",
    Action.NON_SCOPUS: "dim",
    Action.PREPRINT_ONLY: "dim",
    Action.UNKNOWN: "magenta",
}
ACTION_ORDER = [
    Action.LIKELY_MISSING_LINK,
    Action.CHECK_SCOPUS,
    Action.UNKNOWN,
    Action.OK,
    Action.NON_SCOPUS,
    Action.PREPRINT_ONLY,
]


def scopus_label(src: ScopusSourceStatus, art: ScopusArticleStatus) -> str:
    if art == ScopusArticleStatus.SCOPUS_CONFIRMED:
        return "yes"
    if art == ScopusArticleStatus.SCOPUS_NOT_FOUND:
        return "not found"
    if src in (ScopusSourceStatus.COVERED, ScopusSourceStatus.COVERED_TITLE_MATCH):
        return "source" + ("*" if src == ScopusSourceStatus.COVERED_TITLE_MATCH else "")
    if src == ScopusSourceStatus.SERIES_PREVIOUSLY_INDEXED:
        return "series?"
    if src in (ScopusSourceStatus.NOT_IN_LIST, ScopusSourceStatus.OUTSIDE_COVERAGE):
        return "no"
    return "?"


def row_dict(paper: TrackedPaper, c: dict[str, Any]) -> dict[str, Any]:
    w, r, t = c["work"], c["result"], c["target"]
    return {
        "paper": paper.name,
        "cited_title": paper.title,
        "final_doi": paper.journal_doi,
        "preprint_doi": paper.arxiv_doi,
        "arxiv_id": paper.arxiv_id,
        "citing_title": w.title,
        "citing_doi": w.ids.doi,
        "citing_arxiv": w.ids.arxiv,
        "year": w.year,
        "venue": w.venue,
        "issns": ";".join(w.issns),
        "publication_type": r.publication_status.value,
        "found_in": ";".join(c["sources"]),
        "target_version": r.target_version.value,
        "scopus_source_status": r.scopus_source_status.value,
        "scopus_article_status": r.scopus_article_status.value,
        "scopus_eid": w.ids.scopus_eid,
        "linkage_status": r.linkage_status.value,
        "action": r.action.value,
        "confidence": r.confidence,
        "needs_review": r.needs_review,
        "target_evidence": " | ".join(r.target_evidence),
        "source_evidence": " | ".join(r.source_evidence),
        "confidence_evidence": " | ".join(
            f"{e.signal}({e.weight:+.2f})" for e in r.confidence_evidence
        ),
        "matched_reference": t.matched_reference,
        "review_reasons": " | ".join(r.review_reasons),
        "openalex_id": w.ids.openalex,
        "s2_id": w.ids.s2,
        "pmid": w.ids.pmid,
    }


def summary(paper: TrackedPaper, cites: list[dict[str, Any]], db: Database) -> dict[str, Any]:
    results = [c["result"] for c in cites]
    targets = Counter(r.target_version for r in results)
    arts = Counter(r.scopus_article_status for r in results)
    srcs = Counter(r.scopus_source_status for r in results)
    covered = srcs[ScopusSourceStatus.COVERED] + srcs[ScopusSourceStatus.COVERED_TITLE_MATCH]
    discovered = db.conn.execute(
        "SELECT COUNT(*) FROM citation_sources cs JOIN citations c ON c.id = cs.citation_id "
        "WHERE c.paper_id = ?",
        (paper.id,),
    ).fetchone()[0]
    return {
        "discovered_citations": discovered,
        "unique_citing_works": len(cites),
        "confirmed_scopus_articles": arts[ScopusArticleStatus.SCOPUS_CONFIRMED],
        "scopus_covered_source_only": sum(
            1
            for r in results
            if r.scopus_article_status != ScopusArticleStatus.SCOPUS_CONFIRMED
            and r.scopus_source_status
            in (ScopusSourceStatus.COVERED, ScopusSourceStatus.COVERED_TITLE_MATCH)
        ),
        "not_scopus": sum(1 for r in results if r.action == Action.NON_SCOPUS),
        "conference_series_previously_indexed": srcs[ScopusSourceStatus.SERIES_PREVIOUSLY_INDEXED],
        "unknown_coverage": srcs[ScopusSourceStatus.NO_SOURCE_ID]
        + srcs[ScopusSourceStatus.LIST_NOT_LOADED],
        "preprints": sum(1 for r in results if r.action == Action.PREPRINT_ONLY),
        "likely_citing_preprint": targets[TargetVersion.PREPRINT] + targets[TargetVersion.BOTH],
        "likely_citing_version_of_record": targets[TargetVersion.VERSION_OF_RECORD]
        + targets[TargetVersion.BOTH],
        "target_unknown": targets[TargetVersion.UNKNOWN],
        "needs_manual_review": sum(1 for r in results if r.needs_review),
        "actions": {a.value: sum(1 for r in results if r.action == a) for a in Action},
        "scopus_version_of_record": db.get_meta(f"scopus_vor:{paper.id}"),
        "_covered": covered,
    }


def sort_key(c: dict[str, Any]) -> tuple[int, float]:
    return ACTION_ORDER.index(c["result"].action), -c["result"].confidence


def render_report(
    console: Console, db: Database, paper: TrackedPaper, *, verbose: bool = False
) -> None:
    assert paper.id is not None
    cites = sorted(db.citations_for_paper(paper.id), key=sort_key)
    last = db.last_run(paper.id)
    header = Text(paper.name, style="bold")
    sub = []
    if paper.title:
        sub.append(paper.title)
    ids = [
        f"final DOI {paper.journal_doi}" if paper.journal_doi else None,
        f"arXiv {paper.arxiv_id}" if paper.arxiv_id else None,
    ]
    sub.append(" · ".join(i for i in ids if i))
    if last:
        sub.append(f"last sync {last['finished_at'] or last['started_at']} ({last['status']})")
    vor_meta = db.get_meta(f"scopus_vor:{paper.id}")
    if vor_meta and vor_meta.get("citedby_count") is not None:
        # Elsevier asks that displayed counts link back to the Scopus cited-by list.
        sub.append(
            f"Cited {vor_meta['citedby_count']} times in Scopus: "
            f"{vor_meta.get('citedby_url') or vor_meta.get('url')}"
        )
    table = Table(title=None, show_lines=False, expand=True, header_style="bold")
    table.add_column("Citing paper", ratio=5, overflow="fold")
    table.add_column("Year", justify="right", no_wrap=True)
    table.add_column("Venue", ratio=2, overflow="ellipsis", no_wrap=True)
    table.add_column("Type", no_wrap=True)
    table.add_column("Target", no_wrap=True)
    table.add_column("Scopus", no_wrap=True)
    table.add_column("Action", no_wrap=True)
    table.add_column("Conf", justify="right", no_wrap=True)
    for c in cites:
        w, r = c["work"], c["result"]
        title = w.title or "<untitled>"
        if len(title) > 90 and not verbose:
            title = title[:87] + "..."
        label = Text(title)
        if w.ids.doi:
            label.append(f"\n{w.ids.doi}", style="dim")
        if r.needs_review:
            label.append(" ⚑", style="magenta")
        table.add_row(
            label,
            str(w.year or ""),
            w.venue or "",
            r.publication_status.value.replace("_", " "),
            TARGET_LABEL[r.target_version],
            scopus_label(r.scopus_source_status, r.scopus_article_status),
            Text(r.action.value, style=ACTION_STYLE[r.action]),
            f"{r.confidence:.2f}",
        )
    s = summary(paper, cites, db)
    lines = [
        f"Discovered citations: {s['discovered_citations']}  (provider records)",
        f"Unique citing works: {s['unique_citing_works']}",
        f"Confirmed Scopus articles: {s['confirmed_scopus_articles']}",
        f"Scopus-covered source only: {s['scopus_covered_source_only']}",
        f"Not Scopus: {s['not_scopus']}",
        f"Preprint-only citing works: {s['preprints']}",
        f"Conference series previously indexed (heuristic): "
        f"{s['conference_series_previously_indexed']}",
        f"Unknown coverage: {s['unknown_coverage']}",
        f"Likely citing preprint: {s['likely_citing_preprint']}",
        f"Likely citing Version of Record: {s['likely_citing_version_of_record']}",
        f"Cited version unknown: {s['target_unknown']}",
        f"Needs manual review: {s['needs_manual_review']}",
        *(
            [f"Scopus count check: {s['scopus_version_of_record']['count_check']}"]
            if (s["scopus_version_of_record"] or {}).get("count_check")
            else []
        ),
        "Actions: " + ", ".join(f"{k}={v}" for k, v in s["actions"].items() if v),
    ]
    console.print(
        Panel(
            Group(Text("\n".join(x for x in sub if x), style="dim"), table),
            title=header,
            title_align="left",
        )
    )
    console.print(Panel("\n".join(lines), title="Summary", title_align="left"))
    console.print(
        Text(
            "Scopus column: yes = article confirmed via Scopus API; source = journal in "
            "Scopus Source List (article not verified); source* = matched by title "
            "only; series? = earlier editions of the conference indexed (heuristic); "
            "no = not covered; ? = unknown.  ⚑ = needs manual review.",
            style="dim",
        )
    )
    if verbose:
        for c in cites:
            render_details(console, c)


def render_details(console: Console, c: dict[str, Any]) -> None:
    w, r = c["work"], c["result"]
    lines = [
        f"[b]{w.title}[/b]",
        f"DOI: {w.ids.doi or '-'}   arXiv: {w.ids.arxiv or '-'}   EID: {w.ids.scopus_eid or '-'}",
        f"Found in: {', '.join(c['sources'])}",
        f"Action: {r.action.value}  confidence {r.confidence:.2f}",
        "[u]Cited version[/u]: " + r.target_version.value,
        *[f"  - {e}" for e in r.target_evidence],
        "[u]Scopus[/u]: " + f"{r.scopus_source_status.value} / {r.scopus_article_status.value} / "
        f"{r.linkage_status.value}",
        *[f"  - {e}" for e in r.source_evidence],
        *([f"  - {r.scopus_detail}"] if r.scopus_detail else []),
        "[u]Confidence evidence[/u]",
        *[f"  {e.weight:+.2f} {e.signal}: {e.detail}" for e in r.confidence_evidence],
    ]
    if r.review_reasons:
        lines += ["[u]Manual review[/u]", *[f"  - {x}" for x in r.review_reasons]]
    console.print(Panel("\n".join(lines), expand=True))


def report_json(db: Database, paper: TrackedPaper) -> dict[str, Any]:
    assert paper.id is not None
    cites = sorted(db.citations_for_paper(paper.id), key=sort_key)
    s = summary(paper, cites, db)
    s.pop("_covered", None)
    return {
        "paper": paper.model_dump(),
        "summary": s,
        "citations": [
            {
                **row_dict(paper, c),
                "target_evidence": c["result"].target_evidence,
                "source_evidence": c["result"].source_evidence,
                "confidence_evidence": [e.model_dump() for e in c["result"].confidence_evidence],
                "review_reasons": c["result"].review_reasons,
                "discovered_against": c["against"],
            }
            for c in cites
        ],
    }


def export_rows(db: Database, papers: list[TrackedPaper]) -> list[dict[str, Any]]:
    rows = []
    for p in papers:
        assert p.id is not None
        rows += [row_dict(p, c) for c in sorted(db.citations_for_paper(p.id), key=sort_key)]
    return rows


def write_csv(rows: list[dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0].keys()) if rows else ["paper", "citing_title", "citing_doi", "action"]
    with path.open("w", newline="", encoding="utf-8") as fh:
        wr = csv.DictWriter(fh, fieldnames=fields)
        wr.writeheader()
        wr.writerows(rows)


# --------------------------------------------------------------------------- support pack

SUPPORT_ACTIONS = (Action.LIKELY_MISSING_LINK, Action.CHECK_SCOPUS)
SUPPORT_TARGETS = (TargetVersion.PREPRINT, TargetVersion.BOTH)


def support_candidates(db: Database, paper: TrackedPaper) -> list[dict[str, Any]]:
    assert paper.id is not None
    return [
        c
        for c in sorted(db.citations_for_paper(paper.id), key=sort_key)
        # Proven by Scopus (any cited version), or a preprint citation still to verify.
        if (
            c["result"].action == Action.LIKELY_MISSING_LINK
            or (
                c["result"].action in SUPPORT_ACTIONS
                and c["result"].target_version in SUPPORT_TARGETS
            )
        )
        # The Scopus API looked the article up and it is not indexed: nothing to relink.
        and c["result"].scopus_article_status != ScopusArticleStatus.SCOPUS_NOT_FOUND
    ]


# Excel attachment for Scopus support. The first four columns are the Scopus support form's
# fields, verbatim and in its order; the rest support the request. Internal audit data
# (statuses, confidence, coverage evidence) stays in the CSV and Markdown reports.
LINKING_COLUMNS = [
    "Cited article title",
    "Cited article link in Scopus",
    "Citing article",
    "Citing article link in Scopus",
    "Final DOI",
    "Preprint / arXiv",
    "Reference / evidence",
    "Requested correction",
]


def requested_correction(paper: TrackedPaper) -> str:
    vor_eid = (paper.version_ids.get("vor", {}).get("scopus") or "").split("|")[0]
    if vor_eid:
        return f"Link this reference to Scopus record {vor_eid} (published Version of Record)."
    return (
        "Link this reference to the Scopus record of the published Version of Record "
        f"(DOI {paper.journal_doi or '-'})."
    )


def cited_scopus_link(
    paper: TrackedPaper, override: str | None = None, vor_meta: dict | None = None
) -> str | None:
    if override:
        return override
    if vor_meta and vor_meta.get("url"):
        return vor_meta["url"]
    from .providers.scopus import scopus_record_url

    eid = (paper.version_ids.get("vor", {}).get("scopus") or "").split("|")[0]
    return scopus_record_url(eid) if eid else None


def citing_scopus_link(c: dict[str, Any]) -> str | None:
    from .providers.scopus import scopus_record_url

    w = c["work"]
    return w.extra.get("scopus_url") or scopus_record_url(w.ids.scopus_eid)


def reference_found(paper: TrackedPaper, c: dict[str, Any]) -> str:
    """One fixed-pattern sentence for Scopus support, then the reference as deposited."""
    t = c["target"]
    ev = " ".join(t.evidence).lower()
    cites_arxiv = paper.arxiv_id and (
        "manual_confirmed" in t.signals or "arxiv" in ev or paper.arxiv_id in ev
    )
    what = f"arXiv:{paper.arxiv_id}" if cites_arxiv else "the preprint version"
    wrong = c["work"].extra.get("scopus_linked_eid")
    vor_eid = (paper.version_ids.get("vor", {}).get("scopus") or "").split("|")[0]
    published = f"published record {vor_eid}" if vor_eid else "the published record"
    if wrong and wrong != vor_eid:
        text = (
            f"Reference cites {what} and is currently linked in Scopus to record {wrong} "
            f"instead of {published}."
        )
    else:
        text = f"Reference cites {what} instead of {published}."
    if t.matched_reference:
        ref = re.sub(r"^\[Scopus reference #\d+\]\s*", "", " ".join(t.matched_reference.split()))
        text += f' Original reference: "{ref}"'
    return text


def write_linking_xlsx(path: Path, rows: list[dict[str, Any]]) -> None:
    """Excel attachment in the layout requested by Scopus support."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Reference linking"
    ws.append(LINKING_COLUMNS)
    for row in rows:
        ws.append([" ".join(str(row[c]).split()) if row[c] else "" for c in LINKING_COLUMNS])
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [45, 40, 55, 40, 28, 18, 70, 45]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    # Scopus links stay plain text, not clickable hyperlinks: Excel pre-fetches a clicked
    # link without the browser's session cookies, Scopus answers with a one-time Elsevier
    # login redirect, and Excel opens that redirect instead of the record. Copied and pasted
    # into a browser, the plain URL opens the record.
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def support_request_text(
    paper: TrackedPaper,
    rows: list[dict[str, Any]],
    cited_link: str | None,
    vor_meta: dict | None = None,
) -> str:
    short_title = (paper.title or paper.name).split(":")[0].strip()
    arxiv = f"arXiv:{paper.arxiv_id}" if paper.arxiv_id else (paper.arxiv_doi or "-")
    placeholder = "[PASTE THE SCOPUS LINK OF THE PUBLISHED ARTICLE]"
    wrong_records = sorted({r["_wrong_eid"] for r in rows if r.get("_wrong_eid")})
    vor_eid = (paper.version_ids.get("vor", {}).get("scopus") or "").split("|")[0]
    target = f"the published article {vor_eid}" if vor_eid else "the published article"
    # Only facts support can check directly: no inferred totals ("at least N"), and no
    # record-merge request, so the ticket stays a plain reference-linking request.
    count_lines: list[str] = []
    if wrong_records:
        count_lines += [
            "",
            "In Scopus, these references are currently linked to "
            + ("record " if len(wrong_records) == 1 else "records ")
            + ", ".join(wrong_records)
            + ", corresponding to the earlier preprint version, instead of "
            + (f"the published article {vor_eid}." if vor_eid else "the published article."),
        ]
    if vor_meta and vor_meta.get("deficit"):
        count_lines += [
            "",
            f"Scopus currently shows {vor_meta['citedby_count']} citation"
            f"{'' if vor_meta['citedby_count'] == 1 else 's'} for the published article:",
            f"{vor_meta.get('citedby_url') or cited_link}",
            "However, additional Scopus-indexed articles cite the same work through its "
            "earlier arXiv version.",
        ]
    intro = (
        "Several Scopus-indexed articles cite the preprint version of my article, but these "
        "references are not linked to the final published Version of Record in Scopus."
    )
    closing = (
        "I would appreciate it if these references could be reviewed and linked to "
        f"{target} so that the citation links are correctly consolidated."
    )
    lines = [
        "SUBJECT",
        f'Missing citation links to published version of "{short_title}"',
        "",
        "=" * 72,
        "YOUR QUESTION — option A (large numbers of corrections: attach "
        "scopus_reference_linking.xlsx)",
        "=" * 72,
        "",
        intro,
        *count_lines,
        "",
        # Labels exactly as in the Scopus support form's "Your question" instructions.
        f"Cited article title: {paper.title or '-'}",
        f"Cited article link in Scopus: {cited_link or placeholder}",
        "",
        "Please find attached an Excel file listing the citing articles. Each row contains the "
        "cited article and its Scopus link, the citing article and its Scopus link, the "
        "published DOI, the preprint cited, the reference evidence and the requested correction.",
        "",
        "The preprint and the published article represent the same work:",
        f"Preprint: {arxiv}",
        f"Final DOI: {paper.journal_doi or '-'}",
        "",
        closing,
        "",
        "=" * 72,
        "YOUR QUESTION — option B (recommended for a few corrections: paste into the text "
        "box; the Excel file can be attached as well)",
        "=" * 72,
        "",
        intro,
        *count_lines,
        "",
        "The preprint and the published article represent the same work:",
        f"Preprint: {arxiv}",
        f"Final DOI: {paper.journal_doi or '-'}",
        "",
        "Please review the following references and link them to the final published Scopus "
        "record where appropriate.",
        "",
        # Labels exactly as in the Scopus support form's "Your question" instructions.
        f"Cited article title: {paper.title or '-'}",
        f"Cited article link in Scopus: {cited_link or placeholder}",
        "",
    ]
    for row in rows:
        lines += [
            f"Citing article: {row['Citing article']}",
            "Citing article link in Scopus: "
            + (row["Citing article link in Scopus"] or "[LINK SCOPUS]"),
            "",
        ]
    lines += [
        "In each case, the bibliographic reference points to the earlier arXiv version of the "
        "same work rather than to the final Version of Record.",
        closing,
        "",
    ]
    return "\n".join(lines)


def export_support(
    db: Database, paper: TrackedPaper, out_dir: Path, cited_link_override: str | None = None
) -> tuple[list[Path], int, int]:
    """Write the Scopus-support pack. Returns (files, n_candidates, n_missing_scopus_links)."""
    cands = support_candidates(db, paper)
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_path = out_dir / "scopus_missing_citations.csv"
    md_path = out_dir / "scopus_missing_citations.md"
    xlsx_path = out_dir / "scopus_reference_linking.xlsx"
    txt_path = out_dir / "scopus_support_request.txt"
    vor_meta = db.get_meta(f"scopus_vor:{paper.id}")
    cited_link = cited_scopus_link(paper, cited_link_override, vor_meta)
    cited_eid = (paper.version_ids.get("vor", {}).get("scopus") or "").split("|")[0] or None
    arxiv = f"arXiv:{paper.arxiv_id}" if paper.arxiv_id else (paper.arxiv_doi or "")

    rows, linking = [], []
    for c in cands:
        w, r, t = c["work"], c["result"], c["target"]
        link = citing_scopus_link(c)
        rows.append(
            {
                "cited_work_title": paper.title,
                "cited_scopus_link": cited_link,
                "final_doi": paper.journal_doi,
                "preprint_doi": paper.arxiv_doi,
                "arxiv_id": paper.arxiv_id,
                "citing_title": w.title,
                "citing_scopus_link": link,
                "citing_doi": w.ids.doi,
                "citing_venue": w.venue,
                "citing_year": w.year,
                "scopus_eid": w.ids.scopus_eid,
                "reference_text": t.matched_reference,
                "preprint_evidence": " | ".join(r.target_evidence),
                "scopus_evidence": " | ".join(
                    [*r.source_evidence, r.scopus_detail] if r.scopus_detail else r.source_evidence
                ),
                "scopus_source_status": r.scopus_source_status.value,
                "scopus_article_status": r.scopus_article_status.value,
                "linkage_status": r.linkage_status.value,
                "action": r.action.value,
                "confidence": r.confidence,
            }
        )
        linking.append(
            {
                "Cited article title": paper.title,
                "Cited article link in Scopus": cited_link,
                "Final DOI": paper.journal_doi,
                "Preprint / arXiv": arxiv,
                "Citing article": w.title,
                "Citing article link in Scopus": link,
                "Reference / evidence": reference_found(paper, c),
                "Requested correction": requested_correction(paper),
                "_wrong_eid": w.extra.get("scopus_linked_eid")
                if w.extra.get("scopus_linked_eid") != cited_eid
                else None,
            }
        )

    fields = list(rows[0].keys()) if rows else ["cited_work_title", "citing_title", "citing_doi"]
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        wr = csv.DictWriter(fh, fieldnames=fields)
        wr.writeheader()
        wr.writerows(rows)
    write_linking_xlsx(xlsx_path, linking)
    txt_path.write_text(
        support_request_text(paper, linking, cited_link, vor_meta), encoding="utf-8"
    )

    md = [
        f"# Possible missing citation links — {paper.name}",
        "",
        f"**Cited work:** {paper.title or '-'}  ",
        f"**Version of Record DOI:** {paper.journal_doi or '-'}  ",
        f"**Cited article in Scopus:** {cited_link or 'not known (no Scopus API key)'}  ",
        *(
            [
                f"**Cited {vor_meta['citedby_count']} times in Scopus:** "
                f"{vor_meta.get('citedby_url') or cited_link}  ",
                *(
                    [f"**Count check:** {vor_meta['count_check']}  "]
                    if vor_meta.get("count_check")
                    else []
                ),
            ]
            if vor_meta and vor_meta.get("citedby_count") is not None
            else []
        ),
        f"**Preprint:** arXiv:{paper.arxiv_id or '-'} (DOI {paper.arxiv_doi or '-'})",
        "",
        "The citing documents below reference the preprint version of this work. Where they "
        "are indexed in Scopus, the citation may need to be linked to the Version of Record.",
        "",
        "> Generated by citereclaim from public scholarly metadata (Crossref, "
        "OpenAlex, Semantic Scholar) and the Scopus Source Title List"
        + (
            " plus the Scopus API"
            if any(
                c["result"].scopus_article_status == ScopusArticleStatus.SCOPUS_CONFIRMED
                for c in cands
            )
            else ""
        )
        + ". "
        "`CHECK_SCOPUS` rows have not been verified at article level; please confirm "
        "before submitting. Nothing has been sent to Elsevier.",
        "",
        f"Candidates: **{len(cands)}**",
        "",
    ]
    for i, (c, row) in enumerate(zip(cands, rows, strict=True), 1):
        r = c["result"]
        md += [
            f"## {i}. {row['citing_title']}",
            "",
            f"- **Citing DOI:** {row['citing_doi'] or '-'}",
            f"- **Venue / year:** {row['citing_venue'] or '-'} ({row['citing_year'] or '-'})",
            f"- **Scopus record:** {row['citing_scopus_link'] or 'not available'}",
            f"- **Status:** `{r.action.value}` — source `{r.scopus_source_status.value}`, "
            f"article `{r.scopus_article_status.value}`, linkage `{r.linkage_status.value}`",
            f"- **Confidence:** {r.confidence:.2f}",
            "",
            "**Reference as deposited:**",
            "",
            f"> {row['reference_text'] or '(reference text unavailable)'}",
            "",
            "**Evidence that the preprint is cited:**",
            "",
            *[f"- {e}" for e in r.target_evidence],
            "",
            "**Evidence of Scopus coverage / indexing:**",
            "",
            *[f"- {e}" for e in r.source_evidence],
            *([f"- {r.scopus_detail}"] if r.scopus_detail else []),
            "",
        ]
    md_path.write_text("\n".join(md), encoding="utf-8")
    missing = sum(1 for row in linking if not row["Citing article link in Scopus"])
    missing += 0 if cited_link or not linking else 1
    return [xlsx_path, txt_path, csv_path, md_path], len(cands), missing


def dump_json(data: Any) -> str:
    return json.dumps(data, indent=2, ensure_ascii=False, default=str)
